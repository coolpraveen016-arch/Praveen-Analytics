"""Delivery Transit Analysis (D+0) - Period Comparison.
Upload two report Excels (earlier period, later period) -> download a comparison Excel."""
import io, re, hmac
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

NAME = "Name of the Division"
DRR = "D+0 (Delivered+Returned+Redirected)"
HEAD = ["Received", "Invoiced", "Invoiced %", "D+0 Delivered", "D+0 Delivered %",
        "D+0 Redirected", "D+0 Redirected %", "D+0 Returned", "D+0 Returned %",
        DRR, DRR + " (Received) %", DRR + " (Invoiced) %"]
PCT = {2, 4, 6, 8, 10, 11}              # positions in HEAD that are percentages
COMP_HEAD = ["Received Count Difference", "Invoiced % Increase or Decrease",
             "D+0 Delivered % Increase or Decrease", "D+0 Redirected % Increase or Decrease",
             "D+0 Returned % Increase or Decrease",
             DRR + " (Received) % Increase or Decrease", DRR + " (Invoiced) % Increase or Decrease"]
PERIOD = re.compile(r"\d{1,2}\.\d{1,2}\.\d{4}\s*to\s*\d{1,2}\.\d{1,2}\.\d{4}", re.I)


def to_num(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(re.sub(r"[,\s%]", "", str(v)))
    except ValueError:
        return 0.0


def sd(a, b):
    return a / b if b else 0.0


def read_report(xl, sheet):
    """Reads one report sheet -> (title, {key: (division, [Received, Invoiced, Delivered, Redirected, Returned])})"""
    raw = xl.parse(sheet, header=None)
    hdr = next((i for i in range(min(len(raw), 15))
                if isinstance(raw.iat[i, 0], str) and "division" in raw.iat[i, 0].lower()), None)
    if hdr is None or raw.shape[1] < 6:
        return None
    title = next((x.strip() for x in raw.iloc[:hdr, 0] if isinstance(x, str) and x.strip()), sheet)
    rows = {}
    for i in range(hdr + 1, len(raw)):
        n = raw.iat[i, 0]
        if not isinstance(n, str) or not n.strip():
            if rows:
                break
            continue
        if n.strip().lower().startswith("total"):
            continue
        rows[" ".join(n.lower().split())] = (n.strip(), [to_num(raw.iat[i, c]) for c in range(1, 6)])
    return title, rows


def metrics(c):
    rec, inv, dl, rd, rt = c
    drr = dl + rd + rt
    return [rec, inv, sd(inv, rec), dl, sd(dl, rec), rd, sd(rd, rec), rt, sd(rt, rec),
            drr, sd(drr, rec), sd(drr, inv)]


def compare_rows(r1, r2):
    keys = list(r2) + [k for k in r1 if k not in r2]
    out, t1, t2 = [], [0.0] * 5, [0.0] * 5
    for k in keys:
        name = (r2.get(k) or r1.get(k))[0]
        c1 = r1[k][1] if k in r1 else [0.0] * 5
        c2 = r2[k][1] if k in r2 else [0.0] * 5
        t1 = [a + b for a, b in zip(t1, c1)]
        t2 = [a + b for a, b in zip(t2, c2)]
        out.append((name, c1, c2))
    out.append(("Total", t1, t2))
    res = []
    for name, c1, c2 in out:
        m1, m2 = metrics(c1), metrics(c2)
        comp = [m2[0] - m1[0]] + [m2[i] - m1[i] for i in (2, 4, 6, 8, 10, 11)]
        res.append((name, m1, m2, comp))
    return res


def comparison_title(t1, t2):
    m1, m2 = PERIOD.search(t1), PERIOD.search(t2)
    if m1 and m2:
        base = t2[:m2.start()].rstrip(" -")
        return f"Comparison {base} -{m1.group(0)} vs {m2.group(0)}"
    return f"Comparison: {t1} vs {t2}"


def write_sheet(ws, t1, t2, tc, rows):
    green = PatternFill("solid", fgColor="92D050")
    yellow = PatternFill("solid", fgColor="FFFF00")
    s = Side(style="thin")
    box = Border(left=s, right=s, top=s, bottom=s)
    blocks = [(1, t1, HEAD, 0), (15, t2, HEAD, 1), (29, tc, COMP_HEAD, 2)]
    for col0, title, heads, kind in blocks:
        heads = [NAME] + heads
        ws.merge_cells(start_row=1, start_column=col0, end_row=1, end_column=col0 + len(heads) - 1)
        c = ws.cell(1, col0, title)
        c.fill, c.font = green, Font(name="Calibri", bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        for j, h in enumerate(heads):
            c = ws.cell(2, col0 + j, h)
            c.font, c.border = Font(name="Calibri", bold=True), box
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r, (name, m1, m2, cp) in enumerate(rows, start=3):
            last = r == len(rows) + 2
            for j, v in enumerate([name] + [m1, m2, cp][kind]):
                c = ws.cell(r, col0 + j, v)
                c.border, c.font = box, Font(name="Calibri", bold=last)
                if last:
                    c.fill = yellow
                if j == 0:
                    continue
                if kind == 2:
                    c.number_format = "+#,##0;-#,##0;0" if j == 1 else "+0.00%;-0.00%;0.00%"
                else:
                    c.number_format = "0.00%" if (j - 1) in PCT else "#,##0"
    for col0 in (1, 15, 29):
        ws.column_dimensions[L(col0)].width = 26
    for col in range(2, 37):
        if col not in (14, 15, 28, 29):
            ws.column_dimensions[L(col)].width = 13
    ws.column_dimensions["N"].width = ws.column_dimensions["AB"].width = 3
    ws.row_dimensions[2].height = 80
    ws.freeze_panes = "B3"


def run_comparison(bytes1, bytes2):
    x1, x2 = pd.ExcelFile(io.BytesIO(bytes1)), pd.ExcelFile(io.BytesIO(bytes2))
    wb = Workbook()
    wb.remove(wb.active)
    previews, notes = {}, []
    for sh in x1.sheet_names:
        if sh not in x2.sheet_names:
            notes.append(f"Sheet '{sh}' is not in the later file - skipped.")
            continue
        a, b = read_report(x1, sh), read_report(x2, sh)
        if a is None or b is None:
            notes.append(f"Sheet '{sh}' does not look like a report sheet - skipped.")
            continue
        rows = compare_rows(a[1], b[1])
        write_sheet(wb.create_sheet(sh[:31]), a[0], b[0], comparison_title(a[0], b[0]), rows)
        previews[sh] = pd.DataFrame([[r[0]] + r[3] for r in rows], columns=[NAME] + COMP_HEAD)
    if not previews:
        raise ValueError("No matching report sheets found in the two files.")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), previews, notes


def login(st):
    try:
        users = dict(st.secrets["users"])
    except Exception:
        return True                      # no users configured -> no login screen
    if st.session_state.get("auth"):
        return True
    st.title("Login")
    with st.form("login"):
        u = st.text_input("Username")
        p = st.text_input("Password", type="password")
        if st.form_submit_button("Submit"):
            if u in users and hmac.compare_digest(str(users[u]), p):
                st.session_state["auth"] = True
                st.rerun()
            st.error("Wrong username or password")
    return False


def main():
    import streamlit as st
    st.set_page_config(page_title="D+0 Comparison", layout="wide")
    if not login(st):
        return
    st.title("Delivery Transit Analysis (D+0) - Period Comparison")
    st.info("Upload the two report Excel files you want to compare. The comparison Excel is built "
            "for every sheet that exists in both files (AllProducts, SpeedDoc, Parcels...).")
    c1, c2 = st.columns(2)
    f1 = c1.file_uploader("1) Earlier period report", type=["xlsx", "xlsm"])
    f2 = c2.file_uploader("2) Later period report", type=["xlsx", "xlsm"])
    if f1 and f2:
        try:
            data, previews, notes = run_comparison(f1.getvalue(), f2.getvalue())
        except Exception as e:
            st.error(f"Could not compare these files: {e}")
            return
        for n in notes:
            st.warning(n)
        st.download_button("Download comparison Excel", data, "Comparison_Report.xlsx",
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           type="primary")
        for tab, (sh, df) in zip(st.tabs(list(previews)), previews.items()):
            fmt = {COMP_HEAD[0]: "{:+,.0f}", **{c: "{:+.2%}" for c in COMP_HEAD[1:]}}
            tab.dataframe(df.style.format(fmt), hide_index=True, use_container_width=True)


if __name__ == "__main__":
    main()
